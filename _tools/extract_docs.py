from __future__ import annotations

from pathlib import Path


def extract_docx(root: Path, out: Path, fname: str) -> Path:
    from docx import Document

    p = root / fname
    doc = Document(str(p))
    lines: list[str] = []

    for para in doc.paragraphs:
        t = (para.text or "").rstrip()
        if t:
            lines.append(t)

    if doc.tables:
        lines.append("")
        lines.append("[Tables]")
        lines.append("")
        for ti, table in enumerate(doc.tables):
            lines.append(f"-- table {ti} --")
            for row in table.rows:
                cells = [(c.text or "").replace("\n", " ").strip() for c in row.cells]
                lines.append("\t".join(cells))

    out_file = out / f"{fname}.txt"
    out_file.write_text("\n".join(lines), encoding="utf-8")
    return out_file


def extract_xlsx(root: Path, out: Path, fname: str) -> list[Path]:
    import pandas as pd

    p = root / fname
    xl = pd.ExcelFile(p)
    out_files: list[Path] = []
    for sheet in xl.sheet_names:
        df = xl.parse(sheet_name=sheet, dtype=str)
        df = df.where(df.notna(), "")
        safe_sheet = "".join(ch if ch not in r'<>:"/\\|?*' else "_" for ch in sheet)
        out_file = out / f"{fname}.{safe_sheet}.csv"
        df.to_csv(out_file, index=False, encoding="utf-8-sig")
        out_files.append(out_file)
    return out_files


def main() -> None:
    root = Path(__file__).resolve().parent.parent
    out = root / "_extracted_docs"
    out.mkdir(exist_ok=True)

    docx_files = [
        "点灯检软件通信说明.docx",
        "点灯检软件与ICW通信接口变更设计V1.1.docx",
    ]
    xlsx_files = ["点灯数据库表说明.xlsx"]

    for f in docx_files:
        extract_docx(root, out, f)

    for f in xlsx_files:
        extract_xlsx(root, out, f)

    print(f"Extracted to: {out}")


if __name__ == "__main__":
    main()

